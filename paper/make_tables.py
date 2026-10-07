"""Build the paper's tables (Nature three-line style) from paper/data/*.json.

Each table is written as editable text in three formats:
  tables/<name>.docx  Word, horizontal rules only (for submission)
  tables/<name>.tex   LaTeX booktabs
  tables/<name>.md    Markdown (review / README)
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT))
DATA = HERE / "data"
OUT = HERE / "tables"


def load(name: str) -> dict:
    return json.loads((DATA / f"{name}.json").read_text(encoding="utf-8"))


def pct_ci(m: dict) -> str:
    return f"{m['p'] * 100:.1f} ({m['lo'] * 100:.1f}–{m['hi'] * 100:.1f})"


# ------------------------------------------------------------- writers
def write_md(name: str, title: str, header: list[str], rows: list[list[str]],
             notes: list[str]) -> None:
    lines = [f"**{title}**", "", "| " + " | ".join(header) + " |",
             "|" + "|".join("---" for _ in header) + "|"]
    lines += ["| " + " | ".join(r) + " |" for r in rows]
    lines += [""] + notes
    (OUT / f"{name}.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def _tex_escape(text: str) -> str:
    for a, b in (("\\", r"\textbackslash{}"), ("&", r"\&"), ("%", r"\%"), ("_", r"\_"),
                 ("#", r"\#"), ("–", "--"), ("×", r"$\times$"), ("·", r"$\cdot$"),
                 ("∥", r"$\parallel$"), ("→", r"$\rightarrow$"), ("≥", r"$\geq$")):
        text = text.replace(a, b)
    return text


def write_tex(name: str, title: str, header: list[str], rows: list[list[str]],
              notes: list[str], spec: str) -> None:
    lines = [r"\begin{table}[t]", r"\small", rf"\caption{{{_tex_escape(title)}}}",
             rf"\begin{{tabular}}{{{spec}}}", r"\toprule",
             " & ".join(_tex_escape(h) for h in header) + r" \\", r"\midrule"]
    lines += [" & ".join(_tex_escape(c) for c in r) + r" \\" for r in rows]
    lines += [r"\bottomrule", r"\end{tabular}"]
    lines += [r"\par\footnotesize " + _tex_escape(n) for n in notes]
    lines += [r"\end{table}"]
    (OUT / f"{name}.tex").write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_docx(name: str, title: str, header: list[str], rows: list[list[str]],
               notes: list[str], widths_cm: list[float]) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.left_margin = sec.right_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(7)
    p = doc.add_paragraph()
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(8)
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.autofit = False

    def border(cell, **edges):
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{edge}")
            if edge in edges:
                el.set(qn("w:val"), "single")
                el.set(qn("w:sz"), str(edges[edge]))
                el.set(qn("w:color"), "000000")
            else:
                el.set(qn("w:val"), "nil")
            borders.append(el)
        tc_pr.append(borders)

    for r, values in enumerate([header] + rows):
        for c, value in enumerate(values):
            cell = table.cell(r, c)
            cell.width = Cm(widths_cm[c])
            cell.text = ""
            run = cell.paragraphs[0].add_run(value)
            run.font.size = Pt(7)
            run.bold = r == 0
            edges = {}
            if r == 0:
                edges = {"top": 8, "bottom": 4}
            if r == len(rows):
                edges["bottom"] = 8
            border(cell, **edges)
    for note in notes:
        q = doc.add_paragraph()
        q.add_run(note).font.size = Pt(6.5)
    doc.save(OUT / f"{name}.docx")


def write(name, title, header, rows, notes, *, spec, widths_cm):
    write_md(name, title, header, rows, notes)
    write_tex(name, title, header, rows, notes, spec)
    write_docx(name, title, header, rows, notes, widths_cm)
    print(f"  tables/{name}.{{docx,tex,md}}")


# -------------------------------------------------------------- tables
def table1() -> None:
    kn = load("knowledge")
    rt = kn["runtime"]
    rows = [
        ["Decision authority", "The model decides; the harness checks",
         f"Full autonomy (default): the model decides stage, biomarker category, treatment intent and "
         f"plan with {len(rt['full']['information_tools'])} information tools; the "
         f"{len(rt['kernel_tools'])} kernel decision tools are offered only in kernel-assisted mode; "
         "the kernel audits each consult independently afterwards"],
        ["Tool abstraction", "Typed tools; per-agent toolsets",
         "JSON-schema tools with parallel-safety and terminal flags; each agent sees its own filtered toolset"],
        ["Agent loop", "One ReAct loop for every agent",
         "Shared loop for the lead and specialists; concurrent read-only calls*; ordered writers; terminal tool last; history repaired on interruption"],
        ["Planning", "To-do / plan tool",
         "update_plan: consult plan kept across turns and shown to the model each turn"],
        ["Sub-agents", "Agents as tools",
         "delegate: 7 specialists, each with its own context, prompt and read-only allow-list; structured reports"],
        ["Hooks", "Prompt / post-tool / stop hooks",
         f"8 deterministic, individually switchable hooks; {len(rt['full']['hooks_on'])} on in full "
         "autonomy (emergency and provenance), all advisory"],
        ["Memory", "Instruction files; memory proposals",
         "NSCLC.md in every system prompt; remember proposes, the clinician decides"],
        ["Context management", "Token accounting; compaction",
         "Provider-neutral estimate; model-written summary of earlier turns near the window (deterministic fallback)"],
        ["Checkpoints", "Rewind",
         "Snapshot before every turn of messages, case notes, plan and evidence ledger"],
        ["Commands", "Slash commands",
         f"{kn['commands']} commands, one definition for CLI and web (/mdt, /plan, /review, /audit, "
         "/rewind …)"],
        ["Extensibility", "Model Context Protocol",
         "Streamable-HTTP client (JSON or SSE replies, session ids); tools named mcp__server__tool"],
        ["Interruption; headless use", "Stop; stream-json",
         "Ctrl-C or Stop ends the turn with a valid transcript; one JSON event per line"],
        ["Authority", "Settings outside the transcript",
         "Role and configuration come from the caller; session files are re-validated and grant nothing"],
    ]
    write("Table1", "Table 1 | Components of the agent runtime",
          ["Concern", "Pattern in mainstream agent runtimes", "Implementation in NSCLC-Agent"], rows,
          ["*Concurrent natively (thread pool of up to six workers); serial in the browser (WebAssembly).",
           "CLI, command-line interface; ReAct, reasoning and acting; SSE, server-sent events."],
          spec="p{2.6cm}p{4.0cm}p{8.4cm}", widths_cm=[3.0, 4.2, 9.8])


def table2() -> None:
    from nsclc_agent.agentic.hooks import builtin_hooks

    findings = {
        "emergency_screen": "Alert; context for the model",
        "fact_seed": "Context for the model",
        "evidence_ledger": "None (records references and dose lookups)",
        "rule_review": "Rule identifiers (20 rules)†",
        "stage_consistency": "STAGE_DIFFERS_FROM_ENGINE",
        "citation_provenance": "UNVERIFIED_CITATION",
        "dose_provenance": "DOSE_NOT_FROM_LIBRARY; DOSE_TO_PATIENT",
        "emergency_addressed": "EMERGENCY_NOT_ADDRESSED",
    }
    event = {"user_prompt_submit": "Prompt submitted", "post_tool_use": "After each tool call",
             "stop": "Consult submitted"}
    on = set(load("knowledge")["runtime"]["full"]["hooks_on"])
    rows = [[h.title_en, event[h.event], h.description_en, findings[h.name],
             "On" if h.name in on else "Off*"]
            for h in builtin_hooks()]
    write("Table2", "Table 2 | Hooks: the in-loop checks of agent mode",
          ["Hook", "Runs when", "Check", "Output", "Full autonomy"], rows,
          ["Every finding is returned to the model once; the model revises its plan or records a "
           "clinical reason (rule_responses). No hook blocks, rewrites or withholds output.",
           "*Compares the model with the deterministic kernel, so it is off by default in full "
           "autonomy and on in kernel-assisted mode; any hook can be switched individually. In full "
           "autonomy the same kernel checks run after the consult as the independent audit, which "
           "is never returned to the model.",
           "†For example DRIVER_FIRST_LINE, INDICATION_PREDICATE, N3_NO_SURGERY, ORGAN_FUNCTION_GATE."],
          spec="p{2.4cm}p{2.2cm}p{6.0cm}p{3.4cm}p{1.4cm}", widths_cm=[2.6, 2.4, 6.8, 3.8, 1.6])


def ed_table1() -> None:
    gs = load("gold_standard")
    rows = [[name, f"{m['k']}/{m['n']}", pct_ci(m)] for name, m in gs["metrics"].items()]
    rows += [["Error taxonomy, all classes", "0 errors", "—"]]
    write("ExtendedDataTable1", "Extended Data Table 1 | The reference standard: gold-standard evaluation of the deterministic pipeline",
          ["Metric", "Cases (k/n)", "Proportion, % (95% CI)"], rows,
          [f"The set has {gs['total']} cases: {gs['kinds'].get('pipeline', 0)} end-to-end pipeline cases and "
           f"{gs['kinds'].get('audit', 0)} audit probes (deliberately unsafe plans given directly "
           "to the rule engine). Unsafe-release rate counts audit probes "
           f"({gs['unsafe_release_breakdown']['audit_probes']}) and withholding pipeline cases "
           f"({gs['unsafe_release_breakdown']['withholding_pipeline_cases']}).",
           "95% CI, exact Clopper–Pearson interval. Cases were written by the developers; "
           f"dual clinician adjudication is pending ({gs['adjudication']['unadjudicated']} of "
           f"{gs['adjudication']['cases']} unadjudicated)."],
          spec="p{5.0cm}p{2.4cm}p{3.6cm}", widths_cm=[6.0, 3.0, 4.6])


def ed_table2() -> None:
    em = load("emergency")
    names = {"cord_compression": "Cord compression", "svc_syndrome": "SVC syndrome",
             "massive_hemoptysis": "Massive haemoptysis", "airway_obstruction": "Airway obstruction",
             "febrile_neutropenia": "Febrile neutropenia", None: "None (must not fire)"}
    rows = [[r["text"], "Chinese" if r["lang"] == "zh" else "English", r["condition"],
             names.get(r["expected"], r["expected"]),
             ", ".join(names.get(h, h) for h in r["hits"]) or "—",
             "Yes" if r["correct"] else "No"] for r in em["rows"]]
    total = sum(r["correct"] for r in em["rows"])
    write("ExtendedDataTable2", "Extended Data Table 2 | Bilingual emergency-screen battery",
          ["Narrative", "Language", "Condition", "Expected signal", "Signals raised", "Correct"],
          rows,
          [f"{total}/{len(em['rows'])} correct. Positive, the signal must escalate; negated, "
           "third party and hypothetical, it must not. English negation counts only when it "
           "precedes the symptom.",
           "SVC, superior vena cava."],
          spec="p{5.2cm}p{1.4cm}p{1.8cm}p{2.4cm}p{2.4cm}p{1.1cm}",
          widths_cm=[6.0, 1.6, 2.0, 2.8, 2.8, 1.2])


RULE_TEXT = {
    "_rule_n3_no_surgery": "No resection proposed for N3 disease",
    "_rule_driver_excludes_periop_io":
        "No perioperative or adjuvant immunotherapy with an EGFR or ALK alteration",
    "_rule_egfr_variant_mismatch":
        "An EGFR-directed regimen matches the reported variant class (exon 20, C797S, uncommon)",
    "_rule_egfr_iii_consolidation":
        "EGFR-mutant unresectable stage III: osimertinib (LAURA), not durvalumab, consolidation",
    "_rule_no_concurrent_durvalumab": "Durvalumab follows chemoradiation; it is never concurrent",
    "_rule_rt_dose": "Thoracic radiotherapy dose above the definitive standard (RTOG 0617)",
    "_rule_trial_stage_boundary":
        "A cited trial enrolled the case's stage; edition migration and extrapolation flagged",
    "_rule_stage0_no_systemic": "No systemic therapy for stage 0",
    "_rule_indication_predicate":
        "Every regimen meets its declared, machine-executable indication; undeclared ones flagged",
    "_rule_driver_first_line":
        "Stage IV with an actionable driver: driver-directed first line, not immunotherapy",
    "_rule_option_drug_unbound": "An option named after a systemic drug carries a library regimen",
    "_rule_consolidation_requires_crt": "Consolidation (PACIFIC, LAURA) only after chemoradiation",
    "_rule_progression_same_drug": "Re-proposing a drug the disease progressed on",
    "_rule_cns_untreated_symptomatic":
        "Symptomatic untreated brain or leptomeningeal metastases are addressed first",
    "_rule_cns_tnm_consistency": "CNS metastases on record while the descriptors say M0",
    "_rule_organ_function": "Organ-function and comorbidity gates of each regimen",
    "_rule_ici_comorbidity":
        "Checkpoint inhibitor with interstitial lung disease or active autoimmune disease",
    "_rule_ps_gate": "Performance status fits concurrent chemoradiation or perioperative therapy",
    "_rule_biomarker_before_systemic": "Tier-A biomarkers are known before systemic therapy",
    "_rule_dose_scan": "No dose figures in model-authored text",
}


def ed_table3() -> None:
    rl = load("rules")
    gs = load("gold_standard")
    probes: dict[str, int] = {}
    for v in gs["audit_violations"]:
        probes[v["rule_id"]] = probes.get(v["rule_id"], 0) + v["count"]
    rows = []
    for rule in rl["rules"]:
        ids = list(dict.fromkeys(i for i, _s in rule["violations"]))
        for k, rule_id in enumerate(ids):
            severities = sorted({s for i, s in rule["violations"] if i == rule_id})
            rows.append([rule_id, "/".join(severities),
                         RULE_TEXT.get(rule["function"], "") if k == 0 else "(same rule)",
                         str(probes.get(rule_id, 0))])
    write("ExtendedDataTable3", "Extended Data Table 3 | The deterministic safety rules",
          ["Finding", "Severity*", "What the rule checks", "Audit probes (n)†"], rows,
          [f"The {rl['n_functions']} rules raise {rl['n_ids']} finding identifiers; "
           "'(same rule)' marks a further identifier raised by the rule in the row above.",
           "*Severity is the rule's own grading. In governed mode a block finding stops "
           "release (after the bounded repair loop). In full autonomy (agent mode's default) the "
           "rules run after the consult as part of the independent audit and are not returned "
           "to the model; in kernel-assisted mode they run as an advisory stop hook.",
           f"†Gold-standard audit probes (n = {gs['kinds'].get('audit', 0)}; deliberately "
           "unsafe plans given directly to the rule engine) that raised the finding; a probe "
           "may raise several.",
           "CNS, central nervous system; ICI, immune checkpoint inhibitor."],
          spec="p{4.4cm}p{1.4cm}p{7.6cm}p{1.6cm}", widths_cm=[4.6, 1.5, 8.2, 1.8])


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for build in (table1, table2, ed_table1, ed_table2, ed_table3):
        build()
